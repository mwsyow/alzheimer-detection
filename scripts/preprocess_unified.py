#!/usr/bin/env python3
"""Unified, gated preprocessing for native OASIS-1 and MIRIAD T1 MRI.

The workflow follows docs/oasis1-miriad-unified-preprocessing.md. Generated
medical data stay below data/ and are never intended for Git.
"""

from __future__ import annotations

import argparse
import csv
import getpass
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import nibabel as nib
import numpy as np
from openpyxl import load_workbook
from scipy import ndimage
from scipy.io import loadmat


PIPELINE_VERSION = "5"
CLINICA_IMAGE = "alzheimer-unified-clinica:0.11.3-ants2.6.5-r3"
SYNTHSTRIP_IMAGE = "freesurfer/synthstrip:1.8"
EXPECTED_SHAPE = (169, 208, 179)
EXPECTED_ZOOMS = (1.0, 1.0, 1.0)
CLIP_PERCENTILES = (0.5, 99.5)
OASIS_EXPECTED = 235
MIRIAD_EXPECTED = 69
OASIS_ACQUISITION_OVERRIDES = {
    # mpr-1 passed registration but SynthStrip removed inferior midline brain;
    # mpr-2 passed identical automated and visual QC in the pilot.
    "OAS1_0373_MR1": 2,
}


class GateError(RuntimeError):
    """A correctness gate failed."""


@dataclass(frozen=True)
class Paths:
    repo: Path
    output: Path

    @property
    def manifests(self):
        return self.output / "manifests"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".part")
    temporary.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n")
    temporary.replace(path)


def atomic_csv(path: Path, rows: list[dict], *, delimiter: str = ",") -> None:
    if not rows:
        raise GateError(f"refusing to write empty manifest: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".part")
    with temporary.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="") as stream:
        return list(csv.DictReader(stream))


def safe_bids_id(prefix: str, identifier: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9]+", "", identifier).lower()
    return f"sub-{prefix}{value}"


def oasis_rows(repo: Path) -> list[dict]:
    sheet = repo / "data/oasis_cross-sectional_cdr_cleaned.xlsx"
    workbook = load_workbook(sheet, read_only=True, data_only=True)
    values = workbook.active.iter_rows(values_only=True)
    headers = [str(value) for value in next(values)]
    clinical = [dict(zip(headers, row)) for row in values]
    raw_by_id: dict[str, dict[int, Path]] = {}
    for image in (repo / "data").glob(
        "oasis_cross-sectional_disc*/disc*/OAS1_*_MR1/RAW/*_mpr-*_anon.img"
    ):
        match = re.match(r"(OAS1_\d+_MR1)_mpr-(\d+)_anon\.img$", image.name)
        if match:
            subject, acquisition = match.group(1), int(match.group(2))
            if acquisition in raw_by_id.setdefault(subject, {}):
                raise GateError(f"duplicate OASIS acquisition: {subject} mpr-{acquisition}")
            raw_by_id[subject][acquisition] = image.resolve()
    rows = []
    for record in clinical:
        subject = str(record["ID"])
        acquisition = OASIS_ACQUISITION_OVERRIDES.get(subject, 1)
        source = raw_by_id.get(subject, {}).get(acquisition)
        if source is None or not source.with_suffix(".hdr").is_file():
            raise GateError(f"missing OASIS Analyze pair for {subject} mpr-{acquisition}")
        cdr = float(record["CDR"])
        group = "control" if cdr == 0 else "ad" if cdr >= 1 else "ambiguous"
        rows.append(
            {
                "dataset": "oasis1",
                "subject_id": subject,
                "bids_subject": safe_bids_id("oasis", subject.replace("_MR1", "")),
                "bids_session": "ses-baseline",
                "source_image": str(source),
                "source_header": str(source.with_suffix(".hdr")),
                "source_sha256": sha256(source),
                "source_header_sha256": sha256(source.with_suffix(".hdr")),
                "visit": "1",
                "acquisition": str(acquisition),
                "selection_exception": (
                    "mpr-1-failed-synthstrip-qc" if acquisition != 1 else ""
                ),
                "group": group,
                "label_binary": "0" if group == "control" else "1" if group == "ad" else "",
                "age": record["Age"],
                "sex": str(record["M/F"]).lower(),
                "mmse": record["MMSE"],
                "cdr": record["CDR"],
            }
        )
    if len(rows) != OASIS_EXPECTED:
        raise GateError(f"expected {OASIS_EXPECTED} OASIS subjects, found {len(rows)}")
    return rows


def miriad_rows(repo: Path) -> list[dict]:
    metadata = repo / "data/MIRIAD/metadata"
    images = read_csv(metadata / "images.csv")
    clinical = read_csv(metadata / "clinical_assessments.csv")
    by_subject: dict[str, list[dict]] = {}
    for row in images:
        if int(row["acquisition"]) == 1:
            by_subject.setdefault(row["subject_id"], []).append(row)
    clinical_by_subject: dict[str, list[dict]] = {}
    for row in clinical:
        clinical_by_subject.setdefault(row["subject_id"], []).append(row)
    rows = []
    for subject, candidates in sorted(by_subject.items()):
        selected = min(candidates, key=lambda row: (int(row["visit"]), float(row["age"])))
        source = (repo / "data/MIRIAD" / selected["local_path"]).resolve()
        if not source.is_file() or source.stat().st_size != int(selected["size_bytes"]):
            raise GateError(f"missing or incomplete MIRIAD source: {source}")
        assessments = clinical_by_subject.get(subject, [])
        assessment = min(
            assessments,
            key=lambda row: abs(float(row["age"]) - float(selected["age"])),
        )
        group = "ad" if selected["group"].strip().lower() == "ad" else "control"
        visit = int(selected["visit"])
        rows.append(
            {
                "dataset": "miriad",
                "subject_id": subject,
                "bids_subject": safe_bids_id("miriad", subject.removeprefix("miriad_")),
                "bids_session": f"ses-v{visit:02d}",
                "source_image": str(source),
                "source_header": "",
                "source_sha256": sha256(source),
                "source_header_sha256": "",
                "visit": str(visit),
                "acquisition": "1",
                "selection_exception": "nonbaseline-earliest" if visit != 1 else "",
                "group": group,
                "label_binary": "1" if group == "ad" else "0",
                "age": selected["age"],
                "sex": selected["sex"],
                "mmse": assessment["mmse"],
                "cdr": assessment["cdr"],
            }
        )
    if len(rows) != MIRIAD_EXPECTED:
        raise GateError(f"expected {MIRIAD_EXPECTED} MIRIAD subjects, found {len(rows)}")
    exceptions = [row["subject_id"] for row in rows if row["selection_exception"]]
    if exceptions != ["miriad_256"]:
        raise GateError(f"unexpected MIRIAD baseline exceptions: {exceptions}")
    return rows


def validate_inventory(rows: list[dict]) -> None:
    subjects = [(row["dataset"], row["subject_id"]) for row in rows]
    bids = [row["bids_subject"] for row in rows]
    if len(subjects) != len(set(subjects)) or len(bids) != len(set(bids)):
        raise GateError("inventory contains duplicate subject identifiers")
    if len(rows) != OASIS_EXPECTED + MIRIAD_EXPECTED:
        raise GateError(f"expected 304 inputs, found {len(rows)}")
    for row in rows:
        if row["group"] not in {"control", "ad", "ambiguous"}:
            raise GateError(f"invalid group for {row['subject_id']}")


def inventory(paths: Paths) -> Path:
    rows = oasis_rows(paths.repo) + miriad_rows(paths.repo)
    validate_inventory(rows)
    destination = paths.manifests / "cohort.csv"
    atomic_csv(destination, rows)
    atomic_json(
        paths.manifests / "inventory.json",
        {
            "pipeline_version": PIPELINE_VERSION,
            "manifest_sha256": sha256(destination),
            "subjects": len(rows),
            "oasis1": sum(row["dataset"] == "oasis1" for row in rows),
            "miriad": sum(row["dataset"] == "miriad" for row in rows),
        },
    )
    print(f"inventory passed: {destination} ({len(rows)} subjects)")
    return destination


def select_pilot(rows: list[dict], seed: int = 42) -> list[dict]:
    rng = np.random.default_rng(seed)
    selected = []
    for dataset in ("oasis1", "miriad"):
        for group in ("control", "ad"):
            candidates = [row for row in rows if row["dataset"] == dataset and row["group"] == group]
            indices = np.sort(rng.choice(len(candidates), size=5, replace=False))
            selected.extend(candidates[int(index)] for index in indices)
    return sorted(selected, key=lambda row: (row["dataset"], row["subject_id"]))


def load_volume(path: Path) -> tuple[nib.spatialimages.SpatialImage, np.ndarray]:
    image = nib.load(path)
    data = np.asanyarray(image.dataobj)
    if data.ndim == 4 and data.shape[-1] == 1:
        data = data[..., 0]
    if data.ndim != 3:
        raise GateError(f"expected a 3D volume, got {data.shape}: {path}")
    return image, data


def analyze_sagittal_affine(shape, zooms) -> np.ndarray:
    """Return RAS affine for Analyze orient=2 (sagittal unflipped).

    Analyze 7.5 defines its stored axes as posterior-to-anterior,
    inferior-to-superior, and right-to-left for this orientation.
    """
    spacing = np.asarray(zooms[:3], dtype=float)
    linear = np.array(
        [
            [0.0, 0.0, -spacing[2]],
            [spacing[0], 0.0, 0.0],
            [0.0, spacing[1], 0.0],
        ]
    )
    affine = np.eye(4)
    affine[:3, :3] = linear
    affine[:3, 3] = -(linear @ ((np.asarray(shape[:3], dtype=float) - 1) / 2))
    return affine


def source_affine(image: nib.spatialimages.SpatialImage) -> np.ndarray:
    """Read physical geometry, including OASIS' legacy Analyze orient byte."""
    if isinstance(image, nib.spm2analyze.Spm2AnalyzeImage):
        orient = int(np.asarray(image.header["orient"]).tobytes()[0])
        if orient != 2:
            raise GateError(f"unsupported Analyze orientation code: {orient}")
        return analyze_sagittal_affine(image.shape, image.header.get_zooms())
    return np.asarray(image.affine)


def validate_native(path: Path) -> dict:
    image, data = load_volume(path)
    affine = source_affine(image)
    zooms = np.asarray(image.header.get_zooms()[:3], dtype=float)
    if not np.isfinite(data).all() or not np.isfinite(affine).all():
        raise GateError(f"nonfinite native data or affine: {path}")
    if np.ptp(data) <= 0 or np.count_nonzero(data) == 0:
        raise GateError(f"empty or constant native image: {path}")
    determinant = float(np.linalg.det(affine[:3, :3]))
    fov = zooms * np.asarray(data.shape)
    if abs(determinant) < 1e-6 or np.any(zooms <= 0) or np.any((fov < 100) | (fov > 400)):
        raise GateError(f"implausible native geometry: {path}")
    return {
        "shape": list(data.shape),
        "zooms": zooms.tolist(),
        "orientation": list(nib.aff2axcodes(affine)),
        "affine_determinant": determinant,
        "nonzero_fraction": float(np.count_nonzero(data) / data.size),
    }


def stage_bids(rows: list[dict], workspace: Path) -> list[dict]:
    bids = workspace / "bids"
    bids.mkdir(parents=True, exist_ok=True)
    atomic_json(
        bids / "dataset_description.json",
        {
            "Name": "Harmonized native OASIS-1 and MIRIAD",
            "BIDSVersion": "1.10.0",
            "DatasetType": "raw",
        },
    )
    staged = []
    for row in rows:
        source = Path(row["source_image"])
        native_qc = validate_native(source)
        output = bids / row["bids_subject"] / row["bids_session"] / "anat" / (
            f"{row['bids_subject']}_{row['bids_session']}_T1w.nii.gz"
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        image, source_data = load_volume(source)
        affine = source_affine(image)
        output_current = False
        if output.exists():
            staged_image, staged_data = load_volume(output)
            output_current = np.array_equal(source_data, staged_data) and np.allclose(
                affine, staged_image.affine, atol=1e-5
            )
        if not output_current:
            converted = nib.Nifti1Image(source_data, affine)
            converted.set_qform(affine, code=1)
            converted.set_sform(affine, code=1)
            nib.save(converted, output)
        staged_image, staged_data = load_volume(output)
        if not np.array_equal(source_data, staged_data):
            raise GateError(f"BIDS conversion changed voxel values: {source}")
        if not np.allclose(affine, staged_image.affine, atol=1e-5):
            raise GateError(f"BIDS conversion changed affine: {source}")
        staged.append(
            {
                **row,
                "bids_image": str(output),
                "bids_updated": not output_current,
                "native_qc": json.dumps(native_qc),
            }
        )
    participant_rows = [
        {
            "participant_id": row["bids_subject"],
            "source_dataset": row["dataset"],
            "source_subject_id": row["subject_id"],
            "group": row["group"],
            "age": row["age"],
            "sex": row["sex"],
        }
        for row in rows
    ]
    atomic_csv(bids / "participants.tsv", participant_rows, delimiter="\t")
    return staged


def run_command(command: list[str], log: Path) -> None:
    log.parent.mkdir(parents=True, exist_ok=True)
    with log.open("a") as stream:
        result = subprocess.run(command, stdout=stream, stderr=subprocess.STDOUT)
    if result.returncode:
        lines = log.read_text(errors="replace").splitlines()
        tail = "\n".join(lines[-40:])
        raise GateError(
            f"command failed with exit code {result.returncode}: {' '.join(command)}\n"
            f"log: {log}\n--- log tail ---\n{tail}"
        )


def docker_image_id(image: str) -> str:
    result = subprocess.run(
        ["docker", "image", "inspect", image, "--format", "{{.Id}}"],
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout.strip()


def clinica_command(root_mount: str, n_procs: int) -> list[str]:
    """Build a user-mode Clinica command with a writable process directory."""
    return [
        "docker", "run", "--rm",
        "--user", f"{os.getuid()}:{os.getgid()}",
        "-e", "HOME=/tmp",
        "-e", "ITK_GLOBAL_DEFAULT_NUMBER_OF_THREADS=1",
        "-e", "OMP_NUM_THREADS=1",
        "-v", root_mount,
        "-w", "/workspace",
        CLINICA_IMAGE, "clinica", "run", "t1-linear",
        "/workspace/bids", "/workspace/caps",
        "--n_procs", str(n_procs),
        "--working_directory", "/workspace/work",
    ]


def build_images(paths: Paths) -> None:
    run_command(
        [
            "docker", "build", "--network", "host", "-f", str(paths.repo / "containers/unified-preprocessing.Dockerfile"),
            "-t", CLINICA_IMAGE, str(paths.repo),
        ],
        paths.output / "container-build.log",
    )
    run_command(["docker", "pull", SYNTHSTRIP_IMAGE], paths.output / "container-build.log")
    atomic_json(
        paths.manifests / "containers.json",
        {CLINICA_IMAGE: docker_image_id(CLINICA_IMAGE), SYNTHSTRIP_IMAGE: docker_image_id(SYNTHSTRIP_IMAGE)},
    )


def validate_registered(path: Path) -> dict:
    image, data = load_volume(path)
    zooms = tuple(float(value) for value in image.header.get_zooms()[:3])
    if data.shape != EXPECTED_SHAPE or not np.allclose(zooms, EXPECTED_ZOOMS, atol=1e-4):
        raise GateError(f"unexpected Clinica grid {data.shape}/{zooms}: {path}")
    if not np.isfinite(data).all() or np.ptp(data) <= 0:
        raise GateError(f"invalid registered image: {path}")
    determinant = float(np.linalg.det(image.affine[:3, :3]))
    if not math.isfinite(determinant) or abs(determinant) < 0.5 or abs(determinant) > 2.0:
        raise GateError(f"invalid registered affine determinant {determinant}: {path}")
    orientation = nib.aff2axcodes(image.affine)
    if orientation != ("R", "A", "S"):
        raise GateError(f"unexpected registered orientation {orientation}: {path}")
    return {
        "shape": list(data.shape),
        "zooms": list(zooms),
        "orientation": list(orientation),
        "determinant": determinant,
        "affine": np.asarray(image.affine).tolist(),
    }


def validate_transform(path: Path) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise GateError(f"missing or empty Clinica affine transform: {path}")
    values = None
    try:
        matlab = loadmat(path)
        transforms = [
            np.asarray(value, dtype=float).reshape(-1)
            for key, value in matlab.items()
            if not key.startswith("__") and key != "fixed" and np.asarray(value).size >= 12
        ]
        if len(transforms) == 1:
            values = transforms[0]
    except Exception:
        # SciPy raises version-specific MAT parsing exceptions for text ITK
        # transforms; those are handled by the text fallback below.
        pass
    if values is None:
        text_values = re.findall(
            r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?",
            path.read_text(errors="ignore"),
        )
        if len(text_values) >= 12:
            values = np.asarray([float(value) for value in text_values], dtype=float)
    if values is None or values.size < 12 or not np.isfinite(values[:12]).all():
        raise GateError(f"invalid Clinica affine transform: {path}")
    linear = values[:9].reshape(3, 3)
    determinant = abs(float(np.linalg.det(linear)))
    translation = values[9:12]
    if not 0.1 <= determinant <= 10 or np.any(np.abs(translation) > 500):
        raise GateError(f"implausible Clinica affine transform: {path}")


def validate_mask(image_path: Path, mask_path: Path) -> dict:
    image, data = load_volume(image_path)
    mask_image, mask_data = load_volume(mask_path)
    if data.shape != mask_data.shape or not np.allclose(image.affine, mask_image.affine, atol=1e-5):
        raise GateError(f"image/mask grid mismatch: {image_path}")
    unique = np.unique(mask_data)
    if not set(unique).issubset({0, 1}):
        raise GateError(f"nonbinary mask: {mask_path}")
    labels, count = ndimage.label(mask_data > 0)
    sizes = np.bincount(labels.ravel())[1:]
    voxels = int(np.count_nonzero(mask_data))
    volume = voxels * abs(float(np.linalg.det(mask_image.affine[:3, :3])))
    dominant = float(sizes.max() / sizes.sum()) if count and sizes.sum() else 0.0
    if not 500_000 <= volume <= 2_500_000 or dominant < 0.98:
        raise GateError(f"implausible brain mask volume/components: {mask_path}")
    if np.any(data[mask_data == 0] != 0):
        raise GateError(f"SynthStrip output is nonzero outside mask: {image_path}")
    return {"brain_voxels": voxels, "brain_volume_mm3": volume, "dominant_component": dominant}


def normalize_image(image_path: Path, mask_path: Path, output: Path) -> dict:
    image, data = load_volume(image_path)
    _, mask = load_volume(mask_path)
    inside = data[mask > 0].astype(np.float64)
    low, high = np.percentile(inside, CLIP_PERCENTILES)
    clipped = np.clip(inside, low, high)
    mean, std = float(clipped.mean()), float(clipped.std())
    if not math.isfinite(std) or std <= 1e-8:
        raise GateError(f"near-constant brain intensities: {image_path}")
    normalized = np.zeros(data.shape, dtype=np.float32)
    normalized[mask > 0] = (clipped - mean) / std
    output.parent.mkdir(parents=True, exist_ok=True)
    header = image.header.copy()
    header.set_data_dtype(np.float32)
    nib.save(nib.Nifti1Image(normalized, image.affine, header), output)
    _, checked = load_volume(output)
    foreground = checked[mask > 0]
    if not np.isfinite(checked).all() or np.any(checked[mask == 0] != 0):
        raise GateError(f"invalid normalized output: {output}")
    if abs(float(foreground.mean())) > 1e-4 or abs(float(foreground.std()) - 1) > 1e-3:
        raise GateError(f"normalization moments outside tolerance: {output}")
    return {"clip_low": float(low), "clip_high": float(high), "mean": mean, "std": std}


def montage(image_path: Path, mask_path: Path, output: Path) -> None:
    import matplotlib.pyplot as plt

    _, image = load_volume(image_path)
    _, mask = load_volume(mask_path)
    centers = np.round(np.argwhere(mask > 0).mean(axis=0)).astype(int)
    views = [image[centers[0]], image[:, centers[1]], image[:, :, centers[2]]]
    masks = [mask[centers[0]], mask[:, centers[1]], mask[:, :, centers[2]]]
    figure, axes = plt.subplots(1, 3, figsize=(12, 4))
    for axis, view, overlay, title in zip(axes, views, masks, ("sagittal", "coronal", "axial")):
        axis.imshow(np.rot90(view), cmap="gray", vmin=-3, vmax=3)
        axis.contour(np.rot90(overlay), levels=[0.5], colors="red", linewidths=0.5)
        axis.set_title(title)
        axis.axis("off")
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.tight_layout()
    figure.savefig(output, dpi=120)
    plt.close(figure)


def run_pipeline(paths: Paths, rows: list[dict], name: str, n_procs: int) -> Path:
    workspace = paths.output / name
    staged = stage_bids(rows, workspace)
    root_mount = f"{workspace.resolve()}:/workspace"
    caps = workspace / "caps"
    owned_paths = [path for path in (caps, workspace / "work") if path.exists()]
    if owned_paths:
        run_command(
            [
                "docker", "run", "--rm", "-v", root_mount, CLINICA_IMAGE,
                "chown", "-R", f"{os.getuid()}:{os.getgid()}",
                *(f"/workspace/{path.relative_to(workspace)}" for path in owned_paths),
            ],
            workspace / "logs/ownership.log",
        )
    stale = []
    for row in staged:
        subject_dir = caps / "subjects" / row["bids_subject"] / row["bids_session"] / "t1_linear"
        registered = list(subject_dir.glob("*_desc-Crop_res-1x1x1_T1w.nii.gz"))
        if row["bids_updated"] or len(registered) > 1 or (
            len(registered) == 1
            and registered[0].stat().st_mtime < Path(row["bids_image"]).stat().st_mtime
        ):
            stale.append((row, subject_dir))
    if stale:
        for _, subject_dir in stale:
            if subject_dir.exists():
                shutil.rmtree(subject_dir)
            if subject_dir.exists():
                raise GateError(f"could not invalidate stale derivative: {subject_dir}")
        work = workspace / "work/t1-linear"
        if work.exists():
            shutil.rmtree(work)
        print(f"invalidated stale derivatives for {len(stale)} input(s)")
    existing_registered = list(caps.rglob("*_desc-Crop_res-1x1x1_T1w.nii.gz")) if caps.exists() else []
    if len(existing_registered) < len(rows):
        run_command(
            clinica_command(root_mount, n_procs),
            workspace / "logs/clinica.log",
        )
    qc_rows, failures = [], []
    common_affine = None
    for row in staged:
        try:
            candidates = list((caps / "subjects" / row["bids_subject"] / row["bids_session"]).rglob("*_desc-Crop_res-1x1x1_T1w.nii.gz"))
            if len(candidates) != 1:
                raise GateError(f"expected one Clinica output, found {len(candidates)}")
            registered = candidates[0]
            registration_qc = validate_registered(registered)
            affine = np.asarray(registration_qc["affine"])
            if common_affine is None:
                common_affine = affine
            elif not np.allclose(affine, common_affine, atol=1e-5):
                raise GateError("registered image does not use the common MNI affine")
            transforms = list(registered.parent.glob("*_affine.mat"))
            if len(transforms) != 1:
                raise GateError(f"expected one Clinica affine transform, found {len(transforms)}")
            validate_transform(transforms[0])
            subject_output = workspace / "harmonized" / row["bids_subject"] / row["bids_session"]
            stripped = subject_output / f"{row['bids_subject']}_{row['bids_session']}_desc-brain_T1w.nii.gz"
            mask = subject_output / f"{row['bids_subject']}_{row['bids_session']}_desc-brain_mask.nii.gz"
            source_record = subject_output / f"{row['bids_subject']}_{row['bids_session']}_desc-brain_source.json"
            expected_source = {
                "pipeline_version": PIPELINE_VERSION,
                "registered_sha256": sha256(registered),
            }
            current_source = None
            if source_record.exists():
                try:
                    current_source = json.loads(source_record.read_text())
                except json.JSONDecodeError:
                    pass
            if not stripped.exists() or not mask.exists() or current_source != expected_source:
                subject_output.mkdir(parents=True, exist_ok=True)
                run_command(
                    ["docker", "run", "--rm", "-v", f"{registered.parent.resolve()}:/input:ro", "-v", f"{subject_output.resolve()}:/output", SYNTHSTRIP_IMAGE, "-i", f"/input/{registered.name}", "-o", f"/output/{stripped.name}", "-m", f"/output/{mask.name}"],
                    workspace / "logs/synthstrip.log",
                )
                atomic_json(source_record, expected_source)
            mask_qc = validate_mask(stripped, mask)
            normalized = subject_output / f"{row['bids_subject']}_{row['bids_session']}_desc-harmonized_T1w.nii.gz"
            normalization_qc = normalize_image(stripped, mask, normalized)
            montage_path = workspace / "qc/montages" / f"{row['bids_subject']}_{row['bids_session']}.png"
            montage(normalized, mask, montage_path)
            qc_rows.append({
                "dataset": row["dataset"], "subject_id": row["subject_id"],
                "bids_subject": row["bids_subject"], "status": "passed",
                "registered_image": str(registered), "harmonized_image": str(normalized),
                "mask": str(mask), "montage": str(montage_path),
                "registration_qc": json.dumps(registration_qc),
                "mask_qc": json.dumps(mask_qc), "normalization_qc": json.dumps(normalization_qc),
            })
        except Exception as error:
            failures.append({"dataset": row["dataset"], "subject_id": row["subject_id"], "error": repr(error)})
    if qc_rows:
        atomic_csv(workspace / "qc/qc_metrics.csv", qc_rows)
    if failures:
        atomic_csv(workspace / "qc/failures.csv", failures)
        raise GateError(f"{len(failures)} subjects failed; see {workspace / 'qc/failures.csv'}")
    (workspace / "qc/failures.csv").unlink(missing_ok=True)
    provenance = {
        "pipeline_version": PIPELINE_VERSION,
        "cohort": name,
        "subjects": len(rows),
        "manifest_sha256": manifest_digest(rows),
        "qc_sha256": sha256(workspace / "qc/qc_metrics.csv"),
        "containers": {CLINICA_IMAGE: docker_image_id(CLINICA_IMAGE), SYNTHSTRIP_IMAGE: docker_image_id(SYNTHSTRIP_IMAGE)},
        "clip_percentiles": CLIP_PERCENTILES,
    }
    atomic_json(workspace / "provenance.json", provenance)
    return workspace / "provenance.json"


def manifest_digest(rows: list[dict]) -> str:
    payload = json.dumps(rows, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(payload).hexdigest()


def approve_pilot(paths: Paths, reviewer: str) -> None:
    provenance_path = paths.output / "pilot/provenance.json"
    if not provenance_path.exists():
        raise GateError("pilot has not completed successfully")
    provenance = json.loads(provenance_path.read_text())
    approval = {
        "reviewer": reviewer,
        "approved_at": datetime.now(timezone.utc).isoformat(),
        "pipeline_version": PIPELINE_VERSION,
        "pilot_provenance_sha256": sha256(provenance_path),
        "pilot_manifest_sha256": provenance["manifest_sha256"],
        "pilot_qc_sha256": provenance["qc_sha256"],
    }
    atomic_json(paths.manifests / "pilot-approval.json", approval)
    print(f"pilot approved by {reviewer}")


def require_approval(paths: Paths) -> None:
    approval_path = paths.manifests / "pilot-approval.json"
    provenance_path = paths.output / "pilot/provenance.json"
    if not approval_path.exists() or not provenance_path.exists():
        raise GateError("full run requires a completed and approved pilot")
    approval = json.loads(approval_path.read_text())
    if approval.get("pipeline_version") != PIPELINE_VERSION or approval.get("pilot_provenance_sha256") != sha256(provenance_path):
        raise GateError("pilot approval is stale; review and approve the current pilot")


def verify(paths: Paths, name: str) -> None:
    workspace = paths.output / name
    metrics = read_csv(workspace / "qc/qc_metrics.csv")
    failures = []
    for row in metrics:
        try:
            validate_registered(Path(row["registered_image"]))
            validate_mask(Path(row["harmonized_image"]), Path(row["mask"]))
            _, data = load_volume(Path(row["harmonized_image"]))
            _, mask = load_volume(Path(row["mask"]))
            foreground = data[mask > 0]
            if abs(float(foreground.mean())) > 1e-4 or abs(float(foreground.std()) - 1) > 1e-3:
                raise GateError("normalization moments failed")
        except Exception as error:
            failures.append((row["subject_id"], repr(error)))
    if failures:
        raise GateError(f"verification failed: {failures}")
    print(f"verification passed: {len(metrics)} {name} subjects")


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    result.add_argument("--repo", type=Path, default=Path.cwd())
    result.add_argument("--output", type=Path, default=Path("data/unified_preprocessing"))
    sub = result.add_subparsers(dest="command", required=True)
    sub.add_parser("inventory")
    sub.add_parser("build-images")
    pilot = sub.add_parser("pilot")
    pilot.add_argument("--n-procs", type=int, default=max(1, min(2, os.cpu_count() or 1)))
    approve = sub.add_parser("approve-pilot")
    approve.add_argument("--reviewer", default=getpass.getuser())
    full = sub.add_parser("run")
    full.add_argument("--n-procs", type=int, default=max(1, min(2, os.cpu_count() or 1)))
    verify_parser = sub.add_parser("verify")
    verify_parser.add_argument("--cohort", choices=("pilot", "full"), required=True)
    return result


def main() -> None:
    args = parser().parse_args()
    repo = args.repo.resolve()
    output = args.output if args.output.is_absolute() else repo / args.output
    paths = Paths(repo=repo, output=output.resolve())
    if args.command == "inventory":
        inventory(paths)
    elif args.command == "build-images":
        build_images(paths)
    elif args.command == "pilot":
        manifest = inventory(paths)
        selected = select_pilot(read_csv(manifest))
        atomic_csv(paths.manifests / "pilot.csv", selected)
        run_pipeline(paths, selected, "pilot", args.n_procs)
    elif args.command == "approve-pilot":
        approve_pilot(paths, args.reviewer)
    elif args.command == "run":
        require_approval(paths)
        manifest = inventory(paths)
        run_pipeline(paths, read_csv(manifest), "full", args.n_procs)
    elif args.command == "verify":
        verify(paths, args.cohort)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("interrupted; completed Clinica nodes remain resumable", file=sys.stderr)
        raise SystemExit(130)
    except (GateError, subprocess.CalledProcessError) as error:
        print(f"error: {error}", file=sys.stderr)
        raise SystemExit(1)
